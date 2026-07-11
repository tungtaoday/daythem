import pytest


@pytest.fixture
def cls(auth_client):
    client, _ = auth_client
    k = client.post("/api/v1/classes", json={
        "name": "Toán 9", "subject": "Toán", "grade": "9", "default_fee": 500000,
    }).json()
    return client, k["id"]


def test_bulk_add_students(cls):
    client, class_id = cls
    r = client.post(f"/api/v1/classes/{class_id}/students/bulk", json={
        "names": ["Nguyễn Văn A", "  Trần Thị B  ", "", "Lê Văn C"],
    })
    assert r.status_code == 201
    data = r.json()
    assert data["total"] == 3  # dòng rỗng bị bỏ
    names = [s["name"] for s in data["items"]]
    assert names == ["Nguyễn Văn A", "Trần Thị B", "Lê Văn C"]  # đã trim
    # có thật trong lớp
    listed = client.get(f"/api/v1/classes/{class_id}/students").json()
    assert len(listed) == 3


def test_bulk_add_wrong_class_404(cls):
    client, _ = cls
    r = client.post("/api/v1/classes/khong-ton-tai/students/bulk", json={"names": ["A"]})
    assert r.status_code == 404


def test_ocr_disabled_without_key(cls):
    client, _ = cls
    # Test env không set GEMINI_API_KEY → OCR trả 400 rõ ràng, không crash.
    r = client.post("/api/v1/students/ocr", json={"image_base64": "AAAA", "mime_type": "image/png"})
    assert r.status_code == 400


def test_import_disabled_without_key(cls):
    client, _ = cls
    r = client.post("/api/v1/students/import", json={"file_base64": "AAAA", "filename": "ds.pdf"})
    assert r.status_code == 400


def test_import_unsupported_format(cls, monkeypatch):
    # Bật "key" giả để qua guard rồi kiểm định dạng lạ bị từ chối rõ ràng.
    from daythem.config import settings
    monkeypatch.setattr(settings, "GEMINI_API_KEY", "fake-key-for-routing-test")
    client, _ = cls
    r = client.post("/api/v1/students/import", json={"file_base64": "AAAA", "mime_type": "application/zip", "filename": "x.zip"})
    assert r.status_code == 400
    assert "định dạng" in r.json()["detail"].lower()


def test_docx_xlsx_text_extraction():
    import io
    from daythem.adapters.gemini import _docx_to_text, _xlsx_to_text, _kind
    import docx
    d = docx.Document()
    for n in ["Nguyễn Minh Anh", "Trần Bảo Long"]:
        d.add_paragraph(n)
    b = io.BytesIO(); d.save(b)
    assert "Nguyễn Minh Anh" in _docx_to_text(b.getvalue())

    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["STT", "Họ tên"]); ws.append([1, "Phạm Quỳnh Như"])
    b2 = io.BytesIO(); wb.save(b2)
    assert "Phạm Quỳnh Như" in _xlsx_to_text(b2.getvalue())

    assert _kind("image/png", "") == "media"
    assert _kind("application/pdf", "a.pdf") == "pdf"
    assert _kind("", "ds.docx") == "docx"
    assert _kind("", "ds.xlsx") == "xlsx"
    assert _kind("application/zip", "x.bin") == "unknown"


def test_monthly_wrap(cls):
    client, class_id = cls
    s = client.post(f"/api/v1/classes/{class_id}/students/bulk",
                    json={"names": ["An", "Bình", "Cường"]}).json()["items"]
    # điểm danh 1 buổi tháng 7: 2 có mặt / 3
    client.post(f"/api/v1/classes/{class_id}/attendance", json={
        "session_date": "2026-07-05",
        "records": [
            {"student_id": s[0]["id"], "present": True},
            {"student_id": s[1]["id"], "present": True},
            {"student_id": s[2]["id"], "present": False},
        ],
    })
    # thu 2 người
    for i in (0, 1):
        client.post(f"/api/v1/classes/{class_id}/tuition/payment",
                    json={"student_id": s[i]["id"], "paid": True, "amount": 500000, "month": "2026-07"})

    w = client.get("/api/v1/home/monthly-wrap?month=2026-07").json()
    assert w["collected"] == 1000000
    assert w["sessions"] == 1
    assert w["attendance_pct"] == 67  # 2/3
    assert w["students"] == 3
    assert w["classes"] == 1
